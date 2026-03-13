import streamlit as st
import os
from langchain_groq import ChatGroq
from langchain_community.document_loaders import WebBaseLoader
import fitz  # PyMuPDF
from docx import Document

# --- API Key and LLM Initialization ---

@st.cache_resource
def get_llm():
    """
    Initializes and returns the ChatGroq LLM instance.
    Caches the resource to avoid re-initialization on every script run,
    improving performance.
    """
    groq_api_key = 'gsk.........' # Enter your api key here

    return ChatGroq(
        temperature=0.6,
        groq_api_key=groq_api_key,
        model_name="gemma2-9b-it"
    )

# --- Document Text Extraction ---

def extract_text_from_pdf(file_stream):
    """Extracts text content from an uploaded PDF file stream."""
    try:
        pdf_document = fitz.open(stream=file_stream.read(), filetype="pdf")
        text = "".join(page.get_text() for page in pdf_document)
        pdf_document.close()
        return text.strip()
    except Exception as e:
        st.error(f"Error reading PDF file: {e}")
        return None

def extract_text_from_docx(file_stream):
    """Extracts text content from an uploaded DOCX file stream."""
    try:
        doc = Document(file_stream)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
        return "\n".join(full_text).strip()
    except Exception as e:
        st.error(f"Error reading DOCX file: {e}")
        return None

def get_resume_text(uploaded_file):
    """
    Orchestrates text extraction based on file type.
    Returns the extracted text or None if extraction fails.
    """
    if uploaded_file is None:
        return None

    if uploaded_file.name.endswith('.pdf'):
        return extract_text_from_pdf(uploaded_file)
    elif uploaded_file.name.endswith('.docx'):
        return extract_text_from_docx(uploaded_file)
    else:
        st.warning("Unsupported file format. Please upload a PDF or DOCX file.")
        return None

# --- Web Content Extraction ---

def extract_text_from_url(url):
    """Extracts text content from a given URL using WebBaseLoader."""
    if not url or not url.startswith(('http://', 'https://')):
        st.error("Please enter a valid URL.")
        return None
    try:
        loader = WebBaseLoader(url)
        loader.requests_per_second = 1
        page_data = loader.load()
        if page_data:
            return page_data[0].page_content
        else:
            st.error("Could not retrieve any content from the URL.")
            return None
    except Exception as e:
        st.error(f"Failed to extract text from URL: {e}")
        return None